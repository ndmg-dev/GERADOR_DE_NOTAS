"""Geração do arquivo .docx das Notas Explicativas com papel timbrado.

Os parâmetros de layout abaixo já foram validados em produção contra o
documento de referência e não devem ser alterados.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, Twips

from app.services.notas_builder import Nota, Paragrafo, Tabela

logger = logging.getLogger(__name__)

# Geometria conferida contra o arquivo de papel timbrado de referência
# (Papel timbrado.docx): página 11910 x 16840 twips, arte do cabeçalho com
# 2145 twips de altura e a do rodapé com 1436, ambas sangrando até perto da
# borda. As margens de texto abaixo mantêm o corpo fora da arte:
#   topo   2750 > 2145 (folga de ~605 twips)
#   rodapé 16840 - 1700 = 15140 < 15404, onde a arte do rodapé começa
PAGE_SETUP = {
    "width_twips": 11910,  # A4 do papel timbrado de referência
    "height_twips": 16840,
    "margin_top": 2750,  # espaço para o header timbrado
    "margin_bottom": 1700,
    "margin_left": 1134,
    "margin_right": 1134,
}

# Largura de cada arte do timbrado, conforme o documento de referência. Menor
# que a página: a arte é centralizada, deixando uma pequena sangria lateral.
# A altura é derivada da proporção da imagem enviada — com as artes de
# referência isso reproduz 2145 twips no cabeçalho e 1436 no rodapé.
TIMBRADO_LARGURA_TWIPS = {"header": 11397, "footer": 11477}

TABLE_STYLE = {
    "header_fill": "A6A6A6",
    "total_fill": "A6A6A6",
    "border_size": 6,
    "font": "Arial",
    "font_size_body": 22,  # 11pt (half-points)
    "font_size_table": 18,  # 9pt
}

EMU_POR_TWIP = 635  # 1 twip = 1/1440 pol.; 1 pol. = 914400 EMU

# Caracteres de controle que o XML do .docx não aceita (exceto tab e quebras).
_CONTROLE_INVALIDO = {c: None for c in range(0x20) if c not in (0x09, 0x0A, 0x0D)}


def texto_seguro(valor: Any) -> str:
    """Texto utilizável no XML do documento.

    Remove surrogates soltos e caracteres de controle que fariam o `python-docx`
    falhar ao serializar — podem chegar de OCR de baixa qualidade ou de dados
    revisados enviados pelo cliente.
    """
    texto = str(valor)
    if not texto.isascii():
        # Substitui surrogates soltos, que não sobrevivem à serialização.
        texto = texto.encode("utf-8", "replace").decode("utf-8")
    return texto.translate(_CONTROLE_INVALIDO)


class DocxGeneratorService:
    """Monta o documento Word final a partir das notas estruturadas."""

    def generate(
        self,
        notas: list[Nota],
        empresa: dict[str, Any],
        config: dict[str, Any],
        output_path: str | Path,
    ) -> str:
        doc = Document()
        self._apply_page_setup(doc)
        self._set_default_header(doc, empresa)
        self._set_default_footer(doc, empresa)
        self._add_title_block(doc, empresa, config)

        for nota in notas:
            self._add_nota(doc, nota)

        self._add_signature_block(doc, empresa, config)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info("Documento gerado: %s (%d notas)", output_path, len(notas))
        return str(output_path)

    # --------------------------------------------------------------- Página

    def _apply_page_setup(self, doc: Document) -> None:
        for secao in doc.sections:
            secao.page_width = Twips(PAGE_SETUP["width_twips"])
            secao.page_height = Twips(PAGE_SETUP["height_twips"])
            secao.top_margin = Twips(PAGE_SETUP["margin_top"])
            secao.bottom_margin = Twips(PAGE_SETUP["margin_bottom"])
            secao.left_margin = Twips(PAGE_SETUP["margin_left"])
            secao.right_margin = Twips(PAGE_SETUP["margin_right"])

        estilo = doc.styles["Normal"]
        estilo.font.name = TABLE_STYLE["font"]
        estilo.font.size = Pt(TABLE_STYLE["font_size_body"] / 2)
        rpr = estilo.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for atributo in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(atributo), TABLE_STYLE["font"])

    # -------------------------------------------------------- Header/Footer

    def _set_default_header(self, doc: Document, empresa: dict[str, Any]) -> None:
        caminho = empresa.get("timbrado_header_path")
        if not caminho or not Path(caminho).exists():
            logger.info("Empresa sem imagem de header de timbrado; header vazio")
            return

        secao = doc.sections[0]
        paragrafo = secao.header.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        largura = TIMBRADO_LARGURA_TWIPS["header"]
        self._add_floating_image(
            paragrafo,
            Path(caminho),
            largura_twips=largura,
            offset_x_twips=self._inset_horizontal(largura),
            offset_y_twips=0,
        )

    def _set_default_footer(self, doc: Document, empresa: dict[str, Any]) -> None:
        caminho = empresa.get("timbrado_footer_path")
        if not caminho or not Path(caminho).exists():
            logger.info("Empresa sem imagem de footer de timbrado; footer vazio")
            return

        secao = doc.sections[0]
        paragrafo = secao.footer.paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

        largura = TIMBRADO_LARGURA_TWIPS["footer"]
        altura_twips = self._altura_proporcional(Path(caminho), largura)
        self._add_floating_image(
            paragrafo,
            Path(caminho),
            largura_twips=largura,
            offset_x_twips=self._inset_horizontal(largura),
            offset_y_twips=PAGE_SETUP["height_twips"] - altura_twips,
        )

    @staticmethod
    def _inset_horizontal(largura_twips: int) -> int:
        """Centraliza a arte do timbrado na página."""
        return (PAGE_SETUP["width_twips"] - largura_twips) // 2

    @staticmethod
    def _altura_proporcional(imagem: Path, largura_twips: int) -> int:
        from PIL import Image

        try:
            with Image.open(imagem) as img:
                largura_px, altura_px = img.size
            return int(largura_twips * altura_px / largura_px)
        except Exception:
            logger.warning("Não foi possível ler dimensões de %s", imagem, exc_info=True)
            return int(largura_twips * 0.15)

    def _add_floating_image(
        self,
        paragrafo,
        imagem: Path,
        *,
        largura_twips: int,
        offset_x_twips: int,
        offset_y_twips: int,
    ) -> None:
        """Insere a imagem ancorada à página, atrás do texto e sem quebra de linha.

        Reproduz o comportamento do documento de referência: âncora relativa à
        página, ``behindDoc="1"`` e ``<wp:wrapNone/>`` (TextWrappingType.NONE).
        """
        largura_emu = largura_twips * EMU_POR_TWIP
        altura_emu = self._altura_proporcional(imagem, largura_twips) * EMU_POR_TWIP

        run = paragrafo.add_run()
        run.add_picture(str(imagem), width=Emu(largura_emu), height=Emu(altura_emu))

        inline = run._element.find(qn("w:drawing")).find(qn("wp:inline"))
        anchor = self._converter_para_anchor(
            inline,
            offset_x_emu=offset_x_twips * EMU_POR_TWIP,
            offset_y_emu=offset_y_twips * EMU_POR_TWIP,
        )
        drawing = run._element.find(qn("w:drawing"))
        drawing.remove(inline)
        drawing.append(anchor)

    @staticmethod
    def _converter_para_anchor(
        inline, *, offset_x_emu: int, offset_y_emu: int
    ) -> OxmlElement:
        anchor = OxmlElement("wp:anchor")
        anchor.set("distT", "0")
        anchor.set("distB", "0")
        anchor.set("distL", "0")
        anchor.set("distR", "0")
        anchor.set("simplePos", "0")
        anchor.set("relativeHeight", "0")
        anchor.set("behindDoc", "1")  # behindDocument=True
        anchor.set("locked", "0")
        anchor.set("layoutInCell", "1")
        anchor.set("allowOverlap", "1")

        simple_pos = OxmlElement("wp:simplePos")
        simple_pos.set("x", "0")
        simple_pos.set("y", "0")
        anchor.append(simple_pos)

        pos_h = OxmlElement("wp:positionH")
        pos_h.set("relativeFrom", "page")
        offset_h = OxmlElement("wp:posOffset")
        offset_h.text = str(int(offset_x_emu))
        pos_h.append(offset_h)
        anchor.append(pos_h)

        pos_v = OxmlElement("wp:positionV")
        pos_v.set("relativeFrom", "page")
        offset_v = OxmlElement("wp:posOffset")
        offset_v.text = str(int(offset_y_emu))
        pos_v.append(offset_v)
        anchor.append(pos_v)

        for tag in ("wp:extent", "wp:effectExtent", "wp:docPr", "a:graphic"):
            elemento = inline.find(qn(tag))
            if elemento is not None:
                anchor.append(elemento)

        # TextWrappingType.NONE
        wrap_none = OxmlElement("wp:wrapNone")
        extent = anchor.find(qn("wp:extent"))
        if extent is not None:
            extent.addnext(wrap_none)
        else:
            anchor.append(wrap_none)

        return anchor

    # ---------------------------------------------------------------- Corpo

    def _add_title_block(
        self, doc: Document, empresa: dict[str, Any], config: dict[str, Any]
    ) -> None:
        # Cabeçalho no mesmo padrão do documento de referência.
        for texto, tamanho, negrito in (
            (empresa.get("nome", "").upper(), 24, True),
            ("NOTAS EXPLICATIVAS ÀS DEMONSTRAÇÕES CONTÁBEIS", 24, True),
            (f"Findas em 31 de Dezembro de {config.get('ano', '')}", 22, True),
            ("Valores expressos em Reais (R$)", 20, False),
        ):
            if not texto:
                continue
            paragrafo = doc.add_paragraph()
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragrafo.add_run(texto_seguro(texto))
            run.bold = negrito
            run.font.name = TABLE_STYLE["font"]
            run.font.size = Pt(tamanho / 2)

        doc.add_paragraph()

    def _add_nota(self, doc: Document, nota: Nota) -> None:
        titulo = doc.add_paragraph()
        titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = titulo.add_run(texto_seguro(f"Nota {nota.numero:02d} — {nota.titulo}"))
        run.bold = True
        run.font.name = TABLE_STYLE["font"]
        run.font.size = Pt(TABLE_STYLE["font_size_body"] / 2)

        for bloco in nota.conteudo:
            if isinstance(bloco, Paragrafo):
                self._add_paragrafo(doc, bloco.texto)
            elif isinstance(bloco, Tabela):
                self._add_tabela(doc, bloco)

        doc.add_paragraph()

    def _add_paragrafo(self, doc: Document, texto: str) -> None:
        paragrafo = doc.add_paragraph()
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragrafo.paragraph_format.space_after = Pt(6)
        run = paragrafo.add_run(texto_seguro(texto))
        run.font.name = TABLE_STYLE["font"]
        run.font.size = Pt(TABLE_STYLE["font_size_body"] / 2)

    def _add_tabela(self, doc: Document, tabela: Tabela) -> None:
        if tabela.titulo:
            paragrafo = doc.add_paragraph()
            paragrafo.paragraph_format.space_after = Pt(2)
            run = paragrafo.add_run(texto_seguro(tabela.titulo))
            run.bold = True
            run.font.name = TABLE_STYLE["font"]
            run.font.size = Pt(TABLE_STYLE["font_size_table"] / 2)

        linhas_total = len(tabela.linhas) + 1 + (1 if tabela.total else 0)
        tabela_docx = doc.add_table(rows=linhas_total, cols=len(tabela.colunas))
        tabela_docx.autofit = True
        self._aplicar_bordas(tabela_docx)

        # Cabeçalho
        for coluna, texto in enumerate(tabela.colunas):
            celula = tabela_docx.cell(0, coluna)
            self._preencher_celula(
                celula,
                texto,
                negrito=True,
                fill=TABLE_STYLE["header_fill"],
                alinhar_direita=coluna > 0,
            )

        # Corpo
        for indice, linha in enumerate(tabela.linhas, start=1):
            for coluna, valor in enumerate(linha):
                celula = tabela_docx.cell(indice, coluna)
                self._preencher_celula(celula, valor, alinhar_direita=coluna > 0)

        # Totalizador
        if tabela.total:
            ultima = linhas_total - 1
            for coluna, valor in enumerate(tabela.total):
                celula = tabela_docx.cell(ultima, coluna)
                self._preencher_celula(
                    celula,
                    valor,
                    negrito=True,
                    fill=TABLE_STYLE["total_fill"],
                    alinhar_direita=coluna > 0,
                )

        doc.add_paragraph()

    def _preencher_celula(
        self,
        celula,
        texto: str,
        *,
        negrito: bool = False,
        fill: str | None = None,
        alinhar_direita: bool = False,
    ) -> None:
        paragrafo = celula.paragraphs[0]
        for run_existente in list(paragrafo.runs):
            run_existente._element.getparent().remove(run_existente._element)
        paragrafo.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT if alinhar_direita else WD_ALIGN_PARAGRAPH.LEFT
        )
        paragrafo.paragraph_format.space_after = Pt(0)
        run = paragrafo.add_run(texto_seguro(texto))
        run.bold = negrito
        run.font.name = TABLE_STYLE["font"]
        run.font.size = Pt(TABLE_STYLE["font_size_table"] / 2)

        if fill:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:color"), "auto")
            shading.set(qn("w:fill"), fill)
            celula._tc.get_or_add_tcPr().append(shading)

    @staticmethod
    def _aplicar_bordas(tabela_docx) -> None:
        tbl_pr = tabela_docx._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            elemento = OxmlElement(f"w:{lado}")
            elemento.set(qn("w:val"), "single")
            elemento.set(qn("w:sz"), str(TABLE_STYLE["border_size"]))
            elemento.set(qn("w:space"), "0")
            elemento.set(qn("w:color"), "000000")
            borders.append(elemento)
        tbl_pr.append(borders)

    # ---------------------------------------------------------- Assinaturas

    def _add_signature_block(
        self, doc: Document, empresa: dict[str, Any], config: dict[str, Any]
    ) -> None:
        doc.add_paragraph()

        data = config.get("data_aprovacao")
        if data:
            paragrafo = doc.add_paragraph()
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragrafo.add_run(texto_seguro(data))
            run.font.name = TABLE_STYLE["font"]
            run.font.size = Pt(TABLE_STYLE["font_size_body"] / 2)

        doc.add_paragraph()
        doc.add_paragraph()

        for socio in empresa.get("socios") or []:
            self._add_assinatura(
                doc,
                nome=socio.get("nome", "").upper(),
                cargo=socio.get("cargo") or "REPRESENTANTE LEGAL",
                documento=f"CPF: {socio['cpf']}" if socio.get("cpf") else None,
            )

        if empresa.get("contador_nome"):
            cargo = "CONTADOR"
            if empresa.get("contador_crc"):
                cargo = f"CONTADOR - CRC Nº {empresa['contador_crc']}"
            self._add_assinatura(
                doc,
                nome=empresa["contador_nome"].upper(),
                cargo=cargo,
                documento=(
                    f"CPF: {empresa['contador_cpf']}"
                    if empresa.get("contador_cpf")
                    else None
                ),
            )

    def _add_assinatura(
        self, doc: Document, *, nome: str, cargo: str, documento: str | None
    ) -> None:
        if not nome:
            return

        linha = doc.add_paragraph()
        linha.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = linha.add_run("_" * 45)
        run.font.name = TABLE_STYLE["font"]
        run.font.size = Pt(TABLE_STYLE["font_size_body"] / 2)

        for texto, negrito in ((nome, True), (cargo, False), (documento, False)):
            if not texto:
                continue
            paragrafo = doc.add_paragraph()
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragrafo.paragraph_format.space_after = Pt(0)
            run = paragrafo.add_run(texto_seguro(texto))
            run.bold = negrito
            run.font.name = TABLE_STYLE["font"]
            run.font.size = Pt(TABLE_STYLE["font_size_body"] / 2)

        doc.add_paragraph()
        doc.add_paragraph()
