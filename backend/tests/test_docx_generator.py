"""Testes do gerador de .docx, incluindo layout e papel timbrado."""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.services.docx_generator import PAGE_SETUP, TABLE_STYLE, DocxGeneratorService
from app.services.notas_builder import NotasBuilderService
from app.services.pdf_parser import PdfParserService
from tests.test_notas_builder import CONFIG, EMPRESA
from tests.test_pdf_parser import BALANCO_TEXTO, DRE_TEXTO


# Proporções das artes do papel timbrado de referência (Papel timbrado.docx).
HEADER_PX = (2371, 447)
FOOTER_PX = (2391, 299)


def _png(caminho: Path, largura: int = 1200, altura: int = 200) -> Path:
    from PIL import Image

    Image.new("RGB", (largura, altura), (230, 230, 230)).save(caminho)
    return caminho


@pytest.fixture
def notas():
    parser = PdfParserService()
    exercicio = {
        "ano": 2025,
        "balanco": parser.parse_balanco_text(BALANCO_TEXTO),
        "dre": parser.parse_dre_text(DRE_TEXTO),
    }
    return NotasBuilderService().build_all([exercicio], EMPRESA, CONFIG)


@pytest.fixture
def empresa_com_timbrado(tmp_path):
    empresa = dict(EMPRESA)
    empresa["timbrado_header_path"] = str(_png(tmp_path / "header.png", *HEADER_PX))
    empresa["timbrado_footer_path"] = str(_png(tmp_path / "footer.png", *FOOTER_PX))
    return empresa


@pytest.fixture
def documento(notas, empresa_com_timbrado, tmp_path):
    saida = tmp_path / "notas.docx"
    DocxGeneratorService().generate(notas, empresa_com_timbrado, CONFIG, saida)
    return Document(str(saida)), saida


def test_arquivo_gerado(documento) -> None:
    _, saida = documento
    assert saida.exists()
    assert saida.stat().st_size > 0


def test_page_setup_aplicado(documento) -> None:
    doc, _ = documento
    secao = doc.sections[0]

    assert secao.page_width.twips == PAGE_SETUP["width_twips"]
    assert secao.page_height.twips == PAGE_SETUP["height_twips"]
    assert secao.top_margin.twips == PAGE_SETUP["margin_top"]
    assert secao.bottom_margin.twips == PAGE_SETUP["margin_bottom"]
    assert secao.left_margin.twips == PAGE_SETUP["margin_left"]
    assert secao.right_margin.twips == PAGE_SETUP["margin_right"]


def test_todas_as_notas_presentes(documento, notas) -> None:
    doc, _ = documento
    texto = "\n".join(p.text for p in doc.paragraphs)

    for nota in notas:
        assert f"Nota {nota.numero:02d} — {nota.titulo}" in texto


def test_timbrado_flutuante_atras_do_texto(documento) -> None:
    """Header e footer usam wp:anchor com behindDoc=1 e wrapNone."""
    doc, _ = documento
    secao = doc.sections[0]

    for parte in (secao.header, secao.footer):
        anchors = parte._element.findall(f".//{qn('wp:anchor')}")
        assert len(anchors) == 1, "esperada exatamente uma imagem flutuante"

        anchor = anchors[0]
        assert anchor.get("behindDoc") == "1"
        assert anchor.find(qn("wp:wrapNone")) is not None
        assert anchor.find(qn("wp:positionH")).get("relativeFrom") == "page"
        assert anchor.find(qn("wp:positionV")).get("relativeFrom") == "page"
        assert anchor.find(qn("a:graphic")) is not None


def test_timbrado_reproduz_a_geometria_da_referencia(documento) -> None:
    """As artes saem no tamanho e na posição do Papel timbrado.docx."""
    doc, _ = documento
    secao = doc.sections[0]

    # (largura, altura) em twips no documento de referência.
    esperado = {"header": (11397, 2145), "footer": (11477, 1436)}

    for nome, parte in (("header", secao.header), ("footer", secao.footer)):
        anchor = parte._element.findall(f".//{qn('wp:anchor')}")[0]
        extent = anchor.find(qn("wp:extent"))
        largura = int(extent.get("cx")) // 635
        altura = int(extent.get("cy")) // 635
        largura_ref, altura_ref = esperado[nome]

        assert largura == largura_ref
        # Tolerância de arredondamento pixel → EMU.
        assert abs(altura - altura_ref) <= 5

        # Arte centralizada horizontalmente na página.
        offset_x = int(
            anchor.find(qn("wp:positionH")).find(qn("wp:posOffset")).text
        ) // 635
        assert offset_x == (PAGE_SETUP["width_twips"] - largura_ref) // 2

    # O rodapé é ancorado ao pé da página.
    anchor_footer = secao.footer._element.findall(f".//{qn('wp:anchor')}")[0]
    offset_y = int(
        anchor_footer.find(qn("wp:positionV")).find(qn("wp:posOffset")).text
    ) // 635
    altura_footer = int(anchor_footer.find(qn("wp:extent")).get("cy")) // 635
    assert offset_y == PAGE_SETUP["height_twips"] - altura_footer


def test_texto_nao_colide_com_o_timbrado(documento) -> None:
    """As margens precisam manter o corpo fora da arte do timbrado."""
    doc, _ = documento
    secao = doc.sections[0]

    altura_header = (
        int(secao.header._element.findall(f".//{qn('wp:extent')}")[0].get("cy")) // 635
    )
    altura_footer = (
        int(secao.footer._element.findall(f".//{qn('wp:extent')}")[0].get("cy")) // 635
    )

    assert secao.top_margin.twips > altura_header
    assert (
        secao.page_height.twips - secao.bottom_margin.twips
        < secao.page_height.twips - altura_footer
    )


def test_tabelas_com_preenchimento_e_bordas(documento) -> None:
    doc, _ = documento
    assert doc.tables

    tabela = doc.tables[0]
    shading = tabela.cell(0, 0)._tc.find(
        f"{qn('w:tcPr')}/{qn('w:shd')}"
    )
    assert shading is not None
    assert shading.get(qn("w:fill")) == TABLE_STYLE["header_fill"]

    borders = tabela._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    assert borders.find(qn("w:insideH")).get(qn("w:sz")) == str(
        TABLE_STYLE["border_size"]
    )


def test_fonte_das_tabelas(documento) -> None:
    doc, _ = documento
    run = doc.tables[0].cell(0, 0).paragraphs[0].runs[0]

    assert run.font.name == TABLE_STYLE["font"]
    assert run.font.size.pt == TABLE_STYLE["font_size_table"] / 2


def test_bloco_de_assinaturas(documento) -> None:
    doc, _ = documento
    texto = "\n".join(p.text for p in doc.paragraphs)

    # O modelo de referência assina em caixa alta, com cargo e CPF.
    for socio in EMPRESA["socios"]:
        assert socio["nome"].upper() in texto
        assert f"CPF: {socio['cpf']}" in texto
        assert socio["cargo"] in texto

    assert EMPRESA["contador_nome"].upper() in texto
    assert f"CONTADOR - CRC Nº {EMPRESA['contador_crc']}" in texto
    assert f"CPF: {EMPRESA['contador_cpf']}" in texto


def test_titulo_do_documento(documento) -> None:
    doc, _ = documento
    texto = "\n".join(p.text for p in doc.paragraphs)

    assert "NOTAS EXPLICATIVAS ÀS DEMONSTRAÇÕES CONTÁBEIS" in texto
    assert "Findas em 31 de Dezembro de 2025" in texto
    assert "Valores expressos em Reais (R$)" in texto
    assert EMPRESA["nome"].upper() in texto


def test_texto_seguro_preserva_acentos_e_remove_lixo() -> None:
    from app.services.docx_generator import texto_seguro

    assert texto_seguro("Edifícios e Construções") == "Edifícios e Construções"
    assert texto_seguro("tab\there") == "tab\there"
    # Surrogate solto (OCR ruim ou payload malformado) e caractere de controle.
    assert texto_seguro("Ve\udcadculos") == "Ve?culos"
    assert texto_seguro("quebra\x00controle") == "quebracontrole"


def test_dados_com_caractere_invalido_nao_quebram_a_geracao(
    notas, empresa_com_timbrado, tmp_path
) -> None:
    """Um surrogate solto nos dados não pode derrubar a geração inteira."""
    from app.services.notas_builder import Paragrafo

    notas[0].conteudo.insert(0, Paragrafo(texto="Raz\udcadao social estranha"))

    saida = tmp_path / "com_lixo.docx"
    DocxGeneratorService().generate(notas, empresa_com_timbrado, CONFIG, saida)

    assert saida.exists()
    texto = "\n".join(p.text for p in Document(str(saida)).paragraphs)
    assert "Raz?ao social estranha" in texto


def test_empresa_sem_timbrado_nao_falha(notas, tmp_path) -> None:
    saida = tmp_path / "sem_timbrado.docx"
    DocxGeneratorService().generate(notas, dict(EMPRESA), CONFIG, saida)

    doc = Document(str(saida))
    assert doc.sections[0].header._element.findall(f".//{qn('wp:anchor')}") == []
    assert saida.exists()
